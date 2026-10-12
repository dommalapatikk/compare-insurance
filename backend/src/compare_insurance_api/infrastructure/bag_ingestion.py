"""Reproducible ingestion of pinned BAG source files into Supabase."""

from collections.abc import Iterator, Sequence
import csv
from dataclasses import dataclass
from decimal import Decimal
import re
from pathlib import Path
from typing import Any

import httpx

from compare_insurance_api.infrastructure.bag_normalization import (
    DataNormalizationError,
    normalize_premium_row,
)


SWISS_CANTONS = frozenset(
    {
        "AG", "AI", "AR", "BE", "BL", "BS", "FR", "GE", "GL", "GR",
        "JU", "LU", "NE", "NW", "OW", "SG", "SH", "SO", "SZ", "TG",
        "TI", "UR", "VD", "VS", "ZG", "ZH",
    }
)
CONFLICT_COLUMNS = ",".join(
    (
        "premium_year",
        "insurer_id",
        "plan_id",
        "canton",
        "premium_region",
        "age_category",
        "age_subgroup",
        "deductible_chf",
        "accident_coverage",
    )
)


@dataclass(frozen=True, slots=True)
class IngestionSummary:
    source_rows: int
    eligible_rows: int
    excluded_territorial_rows: int


def clean_insurer_display_name(fragments: Sequence[str]) -> str:
    """Extract one readable insurer name from the register's wrapped text."""

    lines = [
        line.strip()
        for fragment in fragments
        for line in fragment.splitlines()
        if line.strip()
    ]
    if not lines:
        raise DataNormalizationError("insurer register row has no name")

    selected: list[str] = []
    for line in lines:
        if selected and selected[-1].endswith("-"):
            joined_line = line
            if selected[-1][-2:-1].islower() and line[:1].isupper():
                joined_line = line[:1].lower() + line[1:]
            selected[-1] = selected[-1][:-1] + joined_line
        else:
            selected.append(line)
        candidate = " ".join(selected)
        if re.search(r"\b(?:AG|SA|GmbH|Genossenschaft|Verein)$", candidate):
            break

    name = " ".join(selected)
    return re.sub(r"(?<=[A-Za-zÀ-ÿ])-(?=[A-Za-zÀ-ÿ])", "", name)


def load_insurer_names(path: Path) -> dict[str, str]:
    """Read BAG insurer numbers and display names from the official workbook."""

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on optional extra
        raise RuntimeError("Install the backend 'data' extra to read BAG XLSX files") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Zugelassene Krankenversicherer"]
    rows = list(sheet.iter_rows(min_row=3, values_only=True))
    names: dict[str, str] = {}
    for index, row in enumerate(rows):
        insurer_number = row[0]
        if not isinstance(insurer_number, int):
            continue
        fragments: list[str] = []
        for grouped_row in rows[index:]:
            if grouped_row is not row and isinstance(grouped_row[0], int):
                break
            if isinstance(grouped_row[2], str):
                fragments.append(grouped_row[2])
        names[str(insurer_number)] = clean_insurer_display_name(fragments)
    workbook.close()
    return names


def load_model_names(path: Path) -> dict[tuple[str, str], str]:
    """Read English model names, falling back to the official German label."""

    names: dict[tuple[str, str], str] = {}
    with path.open(encoding="utf-8-sig", newline="") as source:
        for row in csv.DictReader(source):
            if row["Kategorie"] != "MOD":
                continue
            key = (row["Versicherer"].strip(), row["Tarif"].strip())
            label = (row["Name_EN"] or row["Name_DE"]).strip()
            if key in names and names[key] != label:
                raise DataNormalizationError(f"conflicting model name for {key}")
            names[key] = label
    return names


def iter_normalized_rows(
    premium_path: Path,
    *,
    insurer_names: dict[str, str],
    model_names: dict[tuple[str, str], str],
) -> Iterator[dict[str, Any]]:
    """Yield JSON-safe rows for ordinary Swiss canton comparisons."""

    with premium_path.open(encoding="utf-8-sig", newline="") as source:
        for row in csv.DictReader(source):
            normalized = normalize_premium_row(
                row,
                insurer_names=insurer_names,
                model_names=model_names,
            )
            if normalized["canton"] not in SWISS_CANTONS:
                continue
            amount = normalized["monthly_premium_chf"]
            if isinstance(amount, Decimal):
                normalized["monthly_premium_chf"] = float(amount)
            yield normalized


def validate_source(
    premium_path: Path,
    *,
    insurer_names: dict[str, str],
    model_names: dict[tuple[str, str], str],
) -> IngestionSummary:
    """Validate every source row before any remote write begins."""

    source_rows = 0
    with premium_path.open(encoding="utf-8-sig", newline="") as source:
        for _ in csv.DictReader(source):
            source_rows += 1
    eligible_rows = sum(
        1
        for _ in iter_normalized_rows(
            premium_path,
            insurer_names=insurer_names,
            model_names=model_names,
        )
    )
    if not source_rows or not eligible_rows:
        raise DataNormalizationError("BAG premium source is empty")
    return IngestionSummary(
        source_rows=source_rows,
        eligible_rows=eligible_rows,
        excluded_territorial_rows=source_rows - eligible_rows,
    )


class SupabaseCatalogueWriter:
    """Small PostgREST writer used only by the explicit ingestion command."""

    def __init__(self, *, url: str, service_role_key: str) -> None:
        self._endpoint = f"{url.rstrip('/')}/rest/v1/premium_offers"
        self._headers = {
            "apikey": service_role_key,
            "Authorization": f"Bearer {service_role_key}",
            "Content-Type": "application/json",
        }

    def replace_year(self, premium_year: int) -> None:
        response = httpx.delete(
            self._endpoint,
            params={"premium_year": f"eq.{premium_year}"},
            headers={**self._headers, "Prefer": "return=minimal"},
            timeout=60,
        )
        response.raise_for_status()

    def upsert(self, rows: list[dict[str, Any]]) -> None:
        response = httpx.post(
            self._endpoint,
            params={"on_conflict": CONFLICT_COLUMNS},
            headers={
                **self._headers,
                "Prefer": "resolution=merge-duplicates,return=minimal",
            },
            json=rows,
            timeout=60,
        )
        response.raise_for_status()


def upload_catalogue(
    rows: Iterator[dict[str, Any]],
    *,
    writer: SupabaseCatalogueWriter,
    batch_size: int = 500,
) -> int:
    """Upload validated rows in restart-safe upsert batches."""

    uploaded = 0
    batch: list[dict[str, Any]] = []
    for row in rows:
        batch.append(row)
        if len(batch) == batch_size:
            writer.upsert(batch)
            uploaded += len(batch)
            batch = []
    if batch:
        writer.upsert(batch)
        uploaded += len(batch)
    return uploaded
