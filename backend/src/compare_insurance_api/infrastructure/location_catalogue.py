"""Build the postcode catalogue from pinned swisstopo and BAG sources."""

from collections.abc import Iterable
import csv
from dataclasses import dataclass
from pathlib import Path
import re

from compare_insurance_api.infrastructure.bag_ingestion import SWISS_CANTONS


@dataclass(frozen=True, slots=True, order=True)
class LocationCatalogueRow:
    postal_code: str
    locality: str
    municipality_id: str
    municipality_name: str
    canton: str
    premium_region: str


def _load_region_rules(
    path: Path,
) -> tuple[dict[str, tuple[str, str, str]], dict[str, str], frozenset[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on optional extra
        raise RuntimeError("Install the backend 'data' extra to read BAG XLSX files") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    region_sheet = workbook["Anhang EDI Ver. über die PR"]
    regions = {
        str(row[1]): (str(row[0]), str(row[2]), str(int(row[3])))
        for row in region_sheet.iter_rows(min_row=15, values_only=True)
        if row[0] and isinstance(row[1], int) and row[3] is not None
    }
    mutation_sheet = workbook["Mutationen"]
    mergers: dict[str, str] = {}
    for row in mutation_sheet.iter_rows(values_only=True):
        if not isinstance(row[0], int) or not isinstance(row[3], str):
            continue
        match = re.match(r"fus (\d+)", row[3].strip())
        if match:
            mergers[str(row[0])] = match.group(1)
    workbook.close()
    return regions, mergers, frozenset(canton for canton, _, _ in regions.values())


def build_location_catalogue(
    postcode_path: Path,
    premium_region_path: Path,
) -> tuple[LocationCatalogueRow, ...]:
    """Join official postcodes to the premium-region rules for 2027."""

    regions, mergers, split_cantons = _load_region_rules(premium_region_path)
    excluded_non_residential_municipalities = {"2391", "5391"}
    rows: set[LocationCatalogueRow] = set()

    with postcode_path.open(encoding="utf-8-sig", newline="") as source:
        for source_row in csv.DictReader(source, delimiter=";"):
            canton = source_row["Kantonskürzel"].strip().upper()
            if canton not in SWISS_CANTONS:
                continue
            municipality_id = source_row["BFS-Nr"].strip()
            municipality_name = source_row["Gemeindename"].strip()
            premium_region = "0"

            if canton in split_cantons:
                pricing_municipality_id = mergers.get(municipality_id, municipality_id)
                region_rule = regions.get(pricing_municipality_id)
                if region_rule is None:
                    if municipality_id in excluded_non_residential_municipalities:
                        continue
                    raise ValueError(
                        "No 2027 BAG premium region for municipality "
                        f"{municipality_id} {municipality_name} ({canton})"
                    )
                rule_canton, rule_name, premium_region = region_rule
                if rule_canton != canton:
                    raise ValueError(
                        f"Canton mismatch for municipality {pricing_municipality_id}"
                    )
                municipality_id = pricing_municipality_id
                municipality_name = rule_name

            rows.add(
                LocationCatalogueRow(
                    postal_code=source_row["PLZ4"].strip(),
                    locality=source_row["Ortschaftsname"].strip(),
                    municipality_id=municipality_id,
                    municipality_name=municipality_name,
                    canton=canton,
                    premium_region=premium_region,
                )
            )

    if not rows:
        raise ValueError("Official postcode catalogue is empty")
    return tuple(sorted(rows))


def write_location_catalogue(
    rows: Iterable[LocationCatalogueRow],
    output_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8", newline="") as output:
        writer = csv.DictWriter(
            output,
            fieldnames=[
                "postal_code",
                "locality",
                "municipality_id",
                "municipality_name",
                "canton",
                "premium_region",
            ],
            lineterminator="\n",
        )
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    "postal_code": row.postal_code,
                    "locality": row.locality,
                    "municipality_id": row.municipality_id,
                    "municipality_name": row.municipality_name,
                    "canton": row.canton,
                    "premium_region": row.premium_region,
                }
            )
